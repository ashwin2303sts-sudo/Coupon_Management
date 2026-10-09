from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from contextlib import nullcontext
from unittest.mock import Mock, patch

from django.conf import settings
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.messages.constants import ERROR, SUCCESS
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory
from django.test import SimpleTestCase
from django.template.loader import get_template, render_to_string
from django.urls import resolve

from .models import Booking, Customer
from .views import (
    _canonical_excel_headers,
    _canonical_booking_status,
    _cancel_booking_and_revoke_coupon,
    _coupon_status_for_travel,
    _filter_passengers_by_name,
    _parse_flat_price_range,
    add_booking,
    calculate_discounted_net,
    download_excel_template,
    excel_upload,
)


class FlatPriceRuleTests(SimpleTestCase):
    @patch("core.views.read")
    def test_matching_range_percentage_applies_at_both_boundaries(self, read_rules):
        read_rules.return_value = [
            {
                "airline": "AirAsia",
                "fare_type": "Economy Class",
                "flat_price": "10000 - 50000",
                "percentage": 10,
                "status": "Active",
            },
            {
                "airline": "AirAsia",
                "fare_type": "Economy Class",
                "flat_price": "",
                "percentage": 3,
                "status": "Active",
            },
        ]

        for amount in (Decimal("10000"), Decimal("50000")):
            with self.subTest(amount=amount):
                net_amount, percentage, _ = calculate_discounted_net(
                    amount, "AirAsia", "Economy Class"
                )
                self.assertEqual(percentage, Decimal("10"))
                self.assertEqual(net_amount, amount * Decimal("0.9"))

    @patch("core.views.read")
    def test_amount_outside_range_uses_unbounded_fallback_rule(self, read_rules):
        read_rules.return_value = [
            {
                "airline": "AirAsia",
                "fare_type": "Economy Class",
                "flat_price": "10000 - 50000",
                "percentage": 10,
                "status": "Active",
            },
            {
                "airline": "AirAsia",
                "fare_type": "Economy Class",
                "flat_price": "",
                "percentage": 3,
                "status": "Active",
            },
        ]

        net_amount, percentage, _ = calculate_discounted_net(
            Decimal("50001"), "AirAsia", "Economy Class"
        )

        self.assertEqual(percentage, Decimal("3"))
        self.assertEqual(net_amount, Decimal("48500.97"))

    def test_flat_price_range_rejects_reversed_bounds(self):
        with self.assertRaises(ValueError):
            _parse_flat_price_range("50000 - 10000")


class PassengerNameSearchTests(SimpleTestCase):
    def test_search_filters_by_passenger_name_case_insensitively(self):
        bookings = [
            {"passenger_name": "Anita Rao"},
            {"passenger_name": "Mohan Das"},
        ]

        self.assertEqual(
            _filter_passengers_by_name(bookings, "ANITA"),
            [bookings[0]],
        )


class UpdatedTemplateTests(SimpleTestCase):
    def test_updated_templates_compile(self):
        for template_name in (
            "bookings.html",
            "client_portal.html",
            "coupons.html",
            "customers.html",
            "excel_upload.html",
            "rules.html",
        ):
            with self.subTest(template=template_name):
                get_template(template_name)


class ExcelUploadMessageTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def make_request(self, file_content=b"workbook"):
        request = self.factory.post(
            "/excel-upload/",
            {"excel_file": SimpleUploadedFile("bookings.xlsx", file_content)},
        )
        request.session = {"user_id": "admin"}
        request._messages = FallbackStorage(request)
        return request

    @patch("openpyxl.load_workbook")
    def test_invalid_excel_shows_red_failure_message(self, load_workbook):
        load_workbook.return_value.active.__getitem__.return_value = [
            SimpleNamespace(value="S PNR"),
        ]

        request = self.make_request()
        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        messages = list(request._messages)
        self.assertEqual(messages[0].message, "Something Went Wrong")
        self.assertEqual(messages[0].level_tag, "error")
        self.assertIn("Missing required Excel columns", messages[1].message)
        self.assertIn("Ticket No", messages[1].message)

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("90.00"), Decimal("10"), {}))
    @patch("core.views.timezone.localdate", return_value=date(2026, 10, 8))
    @patch("core.views.Ledger.objects.bulk_create")
    @patch("core.views.Coupon.objects.bulk_create")
    @patch("core.views.Coupon.objects.filter")
    @patch("core.views.Booking.objects.select_related")
    @patch("core.views.Booking.objects.select_for_update")
    @patch("core.views.Booking.objects.bulk_create")
    @patch(
        "core.views.Customer.objects.in_bulk",
        return_value={"CL-1001": Customer(id="CL-1001", name="Client")},
    )
    @patch("core.views.Booking.objects.exclude")
    @patch("core.views.transaction.atomic", return_value=nullcontext())
    @patch("openpyxl.load_workbook")
    def test_successful_excel_import_shows_green_success_message(
        self,
        load_workbook,
        _atomic,
        empty_existing_ticket_nos,
        _customers,
        _bulk_create,
        _select_for_update,
        imported_bookings_query,
        existing_coupons_query,
        coupon_bulk_create,
        ledger_bulk_create,
        _today,
        _calculate_discount,
    ):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        row = [
            "SPNR78901", "CL-1001", "TKT-1001", "Anita Rao", "MAA-DXB", "Web Booking",
            "Round Trip", "AirAsia", "Economy Class", "6E-XY789", "100",
            "2026-10-07", "2026-10-11", "2026-10-01", "Confirmed",
        ]
        second_row = [
            "SPNR78901", "CL-1001", "TKT-1002", "Mohan Das", "DEL-BOM", "Agent Booking",
            "One Way", "IndiGo", "Business Class", "6E-XY790", "250",
            "2026-10-12", "2026-10-12", "2026-10-02", "Cancelled",
        ]
        workbook = load_workbook.return_value
        empty_existing_ticket_nos.return_value.values_list.return_value = []
        _select_for_update.return_value.filter.return_value = []
        existing_coupons_query.return_value.values_list.return_value = []
        workbook.active.__getitem__.return_value = [
            SimpleNamespace(value=header) for header in headers
        ]
        workbook.active.iter_rows.return_value = [tuple(row), tuple(second_row)]

        def assign_booking_ids(bookings, batch_size):
            for index, booking in enumerate(bookings, start=1):
                booking.pk = index
            imported_bookings_query.return_value.filter.return_value = bookings

        _bulk_create.side_effect = assign_booking_ids
        request = self.make_request()
        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        messages = list(request._messages)
        self.assertEqual(messages[0].message, "Booking Added Successfully")
        self.assertEqual(messages[0].level_tag, "success")
        self.assertEqual(
            messages[1].message,
            "Excel upload complete: 2 booking(s) added and 0 booking(s) updated; "
            "1 coupon(s) earned, 1 released and 0 pending until travel date.",
        )
        created_bookings = _bulk_create.call_args.args[0]
        self.assertEqual(len(created_bookings), 2)
        self.assertEqual(
            [(booking.pnr, booking.passenger_name) for booking in created_bookings],
            [("SPNR78901", "Anita Rao"), ("SPNR78901", "Mohan Das")],
        )
        self.assertEqual(
            [booking.ticket_no for booking in created_bookings],
            ["TKT-1001", "TKT-1002"],
        )
        self.assertEqual(
            [(booking.is_auto, booking.is_manual) for booking in created_bookings],
            [(True, False), (True, False)],
        )
        self.assertEqual(
            [booking.status for booking in created_bookings],
            ["Confirmed", "Cancelled"],
        )
        created_coupons = coupon_bulk_create.call_args.args[0]
        self.assertEqual(
            [(coupon.booking.ticket_no, coupon.status, coupon.is_released) for coupon in created_coupons],
            [("TKT-1001", "AVAILABLE", True)],
        )
        self.assertEqual(
            [
                (entry.booking.ticket_no, entry.type, entry.status)
                for entry in ledger_bulk_create.call_args.args[0]
            ],
            [
                ("TKT-1001", "Coupon Earned", "AVAILABLE"),
            ],
        )

    @patch("openpyxl.load_workbook")
    def test_excel_upload_fails_when_status_column_missing(self, load_workbook):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date",
        ]
        workbook = load_workbook.return_value
        workbook.active.__getitem__.return_value = [
            SimpleNamespace(value=header) for header in headers
        ]
        workbook.active.iter_rows.return_value = []
        request = self.make_request()

        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        messages = list(request._messages)
        self.assertEqual(messages[0].message, "Something Went Wrong")
        self.assertIn("Missing required Excel columns", messages[1].message)
        self.assertIn("Status", messages[1].message)

    @patch("core.views.Booking.objects.exclude")
    @patch("openpyxl.load_workbook")
    def test_excel_upload_rejects_row_with_empty_status(self, load_workbook, _exclude):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        row_empty = [
            "SPNR78901", "CL-1001", "TKT-1001", "Anita Rao", "MAA-DXB", "Web Booking",
            "Round Trip", "AirAsia", "Economy Class", "6E-XY789", "100",
            "2026-10-07", "2026-10-11", "2026-10-01", "",
        ]
        workbook = load_workbook.return_value
        _exclude.return_value.values_list.return_value = []
        workbook.active.__getitem__.return_value = [
            SimpleNamespace(value=header) for header in headers
        ]
        workbook.active.iter_rows.return_value = [tuple(row_empty)]
        request = self.make_request()

        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        messages = list(request._messages)
        self.assertEqual(messages[0].message, "Something Went Wrong")
        self.assertIn("Row 2: missing Status", messages[1].message)

    @patch("core.views.Booking.objects.exclude")
    @patch("openpyxl.load_workbook")
    def test_excel_upload_rejects_row_with_invalid_status_value(self, load_workbook, _exclude):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        row_invalid = [
            "SPNR78901", "CL-1001", "TKT-1001", "Anita Rao", "MAA-DXB", "Web Booking",
            "Round Trip", "AirAsia", "Economy Class", "6E-XY789", "100",
            "2026-10-07", "2026-10-11", "2026-10-01", "Hold",
        ]
        workbook = load_workbook.return_value
        _exclude.return_value.values_list.return_value = []
        workbook.active.__getitem__.return_value = [
            SimpleNamespace(value=header) for header in headers
        ]
        workbook.active.iter_rows.return_value = [tuple(row_invalid)]
        request = self.make_request()

        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        messages = list(request._messages)
        self.assertEqual(messages[0].message, "Something Went Wrong")
        self.assertIn("Row 2: Status must be Confirmed or Cancelled.", messages[1].message)

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("90.00"), Decimal("10"), {}))
    @patch("core.views.Ledger.objects.bulk_create")
    @patch("core.views.Coupon.objects.bulk_create")
    @patch("core.views.Coupon.objects.filter")
    @patch("core.views.Booking.objects.select_related")
    @patch("core.views.Booking.objects.select_for_update")
    @patch("core.views.Booking.objects.exclude")
    @patch("core.views.Customer.objects.in_bulk")
    @patch("core.views.transaction.atomic", return_value=nullcontext())
    @patch("openpyxl.load_workbook")
    def test_excel_upload_updates_existing_booking_status_to_confirmed(
        self,
        load_workbook,
        _atomic,
        in_bulk,
        _exclude,
        _select_for_update,
        imported_bookings_query,
        _existing_coupons_query,
        _coupon_bulk_create,
        _ledger_bulk_create,
        _calculate_discount,
    ):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        row = [
            "SPNR78901", "CL-1001", "TKT-1001", "Anita Rao", "MAA-DXB", "Web Booking",
            "Round Trip", "AirAsia", "Economy Class", "6E-XY789", "100",
            "2026-10-07", "2026-10-11", "2026-10-01", "Confirmed",
        ]
        workbook = load_workbook.return_value
        _exclude.return_value.values_list.return_value = ["TKT-1001"]
        in_bulk.return_value = {"CL-1001": Customer(id="CL-1001", name="Client")}
        existing_booking = Booking(
            id=1,
            ticket_no="TKT-1001",
            status="Cancelled",
            pnr="SPNR78901",
            client=in_bulk.return_value["CL-1001"],
            passenger_name="Anita Rao",
            sector="MAA-DXB",
            travel_date=date(2026, 10, 7),
            end_date=date(2026, 10, 11),
            booked_date=date(2026, 10, 1),
            net_amount=Decimal("90.00"),
            base_amount=Decimal("100.00"),
            discount_percentage=Decimal("10.00"),
            airline="AirAsia",
            fare_type="Economy Class",
            airline_pnr="6E-XY789",
            booking_type="Web Booking",
            travel_type="Round Trip",
            is_auto=True,
            is_manual=False,
        )
        existing_booking.save = Mock()
        _select_for_update.return_value.filter.return_value = [existing_booking]
        imported_bookings_query.return_value.filter.return_value = [existing_booking]
        workbook.active.__getitem__.return_value = [
            SimpleNamespace(value=header) for header in headers
        ]
        workbook.active.iter_rows.return_value = [tuple(row)]
        request = self.make_request()

        response = excel_upload(request)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(existing_booking.status, "Confirmed")
        existing_booking.save.assert_called_once()

    def test_excel_status_is_limited_to_confirmed_and_cancelled(self):
        self.assertEqual(_canonical_booking_status(" confirmed "), "Confirmed")
        self.assertEqual(_canonical_booking_status("CANCELLED"), "Cancelled")
        self.assertIsNone(_canonical_booking_status(""))
        self.assertIsNone(_canonical_booking_status(None))
        self.assertIsNone(_canonical_booking_status("Pending"))

    @patch("core.views.Ledger.objects.create")
    @patch("core.views.Coupon.objects.select_for_update")
    def test_cancelling_booking_expires_coupon_and_reverses_earned_points(
        self, select_for_update, create_ledger
    ):
        coupon = SimpleNamespace(
            amount=Decimal("12.50"),
            redeemed_amount=Decimal("2.50"),
            status="AVAILABLE",
            is_released=True,
            customer_id="CL-1001",
            customer=Customer(id="CL-1001", name="Client"),
            customer_name="Client",
            coupon_code="CPN-1",
            id="CPN-1",
            save=Mock(),
        )
        select_for_update.return_value.filter.return_value.first.return_value = coupon
        booking = SimpleNamespace(
            status="Confirmed",
            pnr="SPNR78901",
            pk=1,
            save=Mock(),
        )

        self.assertTrue(_cancel_booking_and_revoke_coupon(booking))

        self.assertEqual(booking.status, "Cancelled")
        self.assertEqual(coupon.amount, 0)
        self.assertEqual(coupon.redeemed_amount, 0)
        self.assertEqual(coupon.status, "EXPIRED")
        self.assertFalse(coupon.is_released)
        coupon.save.assert_called_once()
        self.assertEqual(create_ledger.call_args.kwargs["type"], "Coupon Reversed")
        self.assertEqual(create_ledger.call_args.kwargs["amount"], Decimal("-15.00"))

    @patch("core.views.Ledger.objects.create")
    @patch("core.views.Coupon.objects.select_for_update")
    def test_expired_coupon_is_not_reversed_twice(self, select_for_update, create_ledger):
        coupon = SimpleNamespace(
            status="EXPIRED",
            amount=Decimal("0"),
            redeemed_amount=Decimal("0"),
        )
        select_for_update.return_value.filter.return_value.first.return_value = coupon
        booking = SimpleNamespace(
            status="Cancelled",
            pnr="SPNR78901",
            save=Mock(),
        )

        self.assertFalse(_cancel_booking_and_revoke_coupon(booking))
        create_ledger.assert_not_called()
        self.assertEqual(booking.status, "Cancelled")

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("90.00"), Decimal("10"), {}))
    @patch("core.views.Booking.objects.create")
    @patch("core.views.Booking.objects.filter")
    @patch("core.views.Customer.objects.filter")
    def test_manually_added_booking_sets_manual_flag_only(
        self, customers, existing_bookings, create_booking, _calculate_discount
    ):
        customer = Customer(id="CL-1001", name="Client")
        customers.return_value.first.return_value = customer
        existing_bookings.return_value.exists.return_value = False

        request = self.factory.post(
            "/bookings/add/",
            {
                "s_pnr": "SPNR78901",
                "ticket_no": "TKT-1003",
                "passenger_name": "Anita Rao",
                "client_id": "CL-1001",
                "sector": "MAA-DXB",
                "booking_type": "Web Booking",
                "travel_type": "One Way",
                "airline": "AirAsia",
                "fare_type": "Economy Class",
                "airline_pnr": "6E-XY789",
                "net_amount": "100",
                "travel_date": "2026-10-07",
                "end_date": "2026-10-11",
                "booked_date": "2026-10-01",
            },
        )
        request.session = {"user_id": "admin"}
        request._messages = FallbackStorage(request)

        response = add_booking(request)

        self.assertEqual(response.status_code, 302)
        booking = create_booking.call_args.kwargs
        self.assertEqual(booking["status"], "Confirmed")
        self.assertFalse(booking["is_auto"])
        self.assertTrue(booking["is_manual"])

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("-800.00"), Decimal("0"), None))
    @patch("core.views.Booking.objects.create")
    @patch("core.views.Booking.objects.filter")
    @patch("core.views.Customer.objects.filter")
    def test_add_booking_accepts_negative_amount_and_sets_status_cancelled(
        self, customers, existing_bookings, create_booking, _calculate_discount
    ):
        customer = Customer(id="CL-1001", name="Client")
        customers.return_value.first.return_value = customer
        existing_bookings.return_value.exists.return_value = False

        request = self.factory.post(
            "/bookings/add/",
            {
                "s_pnr": "SPNRNEG1",
                "ticket_no": "TKT-NEG-001",
                "passenger_name": "Test Passenger",
                "client_id": "CL-1001",
                "sector": "DEL - MAA",
                "booking_type": "Web Booking",
                "travel_type": "One Way",
                "airline": "Air India",
                "fare_type": "Economy Class",
                "airline_pnr": "AIRNEG1",
                "net_amount": "-800.00",
                "travel_date": "2026-10-15",
                "end_date": "2026-10-15",
                "booked_date": "2026-10-09",
                "status": "Confirmed",  # Even if Confirmed is passed, negative amount forces Cancelled
            },
        )
        request.session = {"user_id": "admin"}
        request._messages = FallbackStorage(request)

        response = add_booking(request)
        self.assertEqual(response.status_code, 302)
        booking = create_booking.call_args.kwargs
        self.assertEqual(booking["status"], "Cancelled")
        self.assertEqual(booking["net_amount"], Decimal("-800.00"))
        self.assertEqual(booking["base_amount"], Decimal("-800.00"))

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("-500.00"), Decimal("0"), None))
    @patch("core.views.Ledger.objects.bulk_create")
    @patch("core.views.Coupon.objects.bulk_create")
    @patch("core.views.Coupon.objects.filter")
    @patch("core.views.Booking.objects.select_related")
    @patch("core.views.Booking.objects.select_for_update")
    @patch("core.views.Booking.objects.bulk_create")
    @patch("core.views.Booking.objects.exclude")
    @patch("core.views.Customer.objects.in_bulk")
    @patch("core.views.transaction.atomic", return_value=nullcontext())
    @patch("openpyxl.load_workbook")
    def test_excel_upload_accepts_negative_net_amount_and_sets_status_cancelled(
        self,
        load_workbook,
        _atomic,
        in_bulk,
        _exclude,
        booking_bulk_create,
        _select_for_update,
        imported_bookings_query,
        _existing_coupons_query,
        coupon_bulk_create,
        _ledger_bulk_create,
        _calculate_discount,
    ):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        row = [
            "SPNR-NEG2", "CL-1001", "TKT-NEG-002", "Negative Passenger", "DEL-BOM", "Web Booking",
            "One Way", "Air India", "Economy Class", "AI-NEG2", "-500",
            "2026-10-20", "2026-10-20", "2026-10-09", "Confirmed",
        ]
        workbook = load_workbook.return_value
        _exclude.return_value.values_list.return_value = []
        workbook.active.__getitem__.return_value = [SimpleNamespace(value=h) for h in headers]
        workbook.active.iter_rows.return_value = [tuple(row)]

        customer = Customer(id="CL-1001", name="Rohan Sharma")
        in_bulk.return_value = {"CL-1001": customer}
        _select_for_update.return_value.filter.return_value = []

        imported_booking = Booking(
            ticket_no="TKT-NEG-002",
            net_amount=Decimal("-500.00"),
            status="Cancelled",
        )
        imported_bookings_query.return_value.filter.return_value = [imported_booking]
        _existing_coupons_query.return_value.values_list.return_value = []

        request = self.make_request()
        response = excel_upload(request)
        self.assertEqual(response.status_code, 302)
        created_bookings = booking_bulk_create.call_args.args[0]
        self.assertEqual(created_bookings[0].status, "Cancelled")
        self.assertEqual(created_bookings[0].net_amount, Decimal("-500.00"))
        coupon_bulk_create.assert_not_called()

    @patch("core.views.calculate_discounted_net", return_value=(Decimal("1000.00"), Decimal("0"), None))
    @patch("core.views.Ledger.objects.bulk_create")
    @patch("core.views.Coupon.objects.bulk_create")
    @patch("core.views.Coupon.objects.filter")
    @patch("core.views.Booking.objects.select_related")
    @patch("core.views.Booking.objects.select_for_update")
    @patch("core.views.Booking.objects.bulk_create")
    @patch("core.views.Booking.objects.exclude")
    @patch("core.views.Customer.objects.in_bulk")
    @patch("core.views.transaction.atomic", return_value=nullcontext())
    @patch("openpyxl.load_workbook")
    def test_excel_upload_skips_duplicate_ticket_numbers_keeping_first_occurrence(
        self,
        load_workbook,
        _atomic,
        in_bulk,
        _exclude,
        booking_bulk_create,
        _select_for_update,
        imported_bookings_query,
        _existing_coupons_query,
        coupon_bulk_create,
        _ledger_bulk_create,
        _calculate_discount,
    ):
        headers = [
            "S PNR", "Client ID", "Ticket No", "Passenger Name", "Sector", "Booking Type",
            "Travel Type", "Airline Name", "Airline Type", "Airline PNR",
            "Net Amount", "Travel Start Date", "Travel End Date", "Booking Date", "Status",
        ]
        ticket_numbers = ["101", "102", "103", "104", "104", "104", "105", "106"]
        rows = [
            [
                f"SPNR-{tkt}", "CL-1001", tkt, f"Passenger {tkt}", "DEL-BOM", "Web Booking",
                "One Way", "Air India", "Economy Class", f"AI-{tkt}", "1000",
                "2026-10-20", "2026-10-20", "2026-10-09", "Confirmed",
            ]
            for tkt in ticket_numbers
        ]
        workbook = load_workbook.return_value
        _exclude.return_value.values_list.return_value = []
        workbook.active.__getitem__.return_value = [SimpleNamespace(value=h) for h in headers]
        workbook.active.iter_rows.return_value = [tuple(r) for r in rows]

        customer = Customer(id="CL-1001", name="Rohan Sharma")
        in_bulk.return_value = {"CL-1001": customer}
        _select_for_update.return_value.filter.return_value = []
        imported_bookings_query.return_value.filter.return_value = []
        _existing_coupons_query.return_value.values_list.return_value = []

        request = self.make_request()
        response = excel_upload(request)
        self.assertEqual(response.status_code, 302)
        created_bookings = booking_bulk_create.call_args.args[0]
        imported_tickets = [b.ticket_no for b in created_bookings]
        self.assertEqual(imported_tickets, ["101", "102", "103", "104", "105", "106"])

    def test_django_templates_include_message_context_processor(self):
        processors = settings.TEMPLATES[0]["OPTIONS"]["context_processors"]
        self.assertIn("django.contrib.messages.context_processors.messages", processors)

    def test_excel_upload_page_renders_success_and_error_alerts(self):
        for level, text, css_class in (
            (SUCCESS, "Booking Added Successfully", "success"),
            (ERROR, "Something Went Wrong", "error"),
        ):
            with self.subTest(level=level):
                request = self.factory.get("/excel-upload/")
                request.session = {"user_id": "admin"}
                request.resolver_match = resolve("/excel-upload/")
                request._messages = FallbackStorage(request)
                request._messages.add(level, text)

                html = render_to_string(
                    "base.html",
                    {"current_user": "Administrator"},
                    request=request,
                )

                self.assertIn(f'class="alert {css_class}"', html)
                self.assertIn(text, html)

    def test_excel_upload_accepts_common_header_variations(self):
        self.assertEqual(
            _canonical_excel_headers(
                [
                    "PNR",
                    "Customer ID",
                    "Passenger Name",
                    "Sector",
                    "Booking Type",
                    "Travel Type",
                    "Airline",
                    "Fare Type",
                    "Airline PNR",
                    "Amount",
                    "Travel Date",
                    "End Date",
                    "Booked Date",
                    "Passenger ID",
                ]
            ),
            [
                "S PNR",
                "Client ID",
                "Passenger Name",
                "Sector",
                "Booking Type",
                "Travel Type",
                "Airline Name",
                "Airline Type",
                "Airline PNR",
                "Net Amount",
                "Travel Start Date",
                "Travel End Date",
                "Booking Date",
                "Passenger ID",
            ],
        )

    def test_downloaded_excel_contains_only_the_upload_page_headers(self):
        from io import BytesIO
        from openpyxl import load_workbook

        request = self.factory.get("/excel-upload/template/")
        request.session = {"user_id": "admin"}

        response = download_excel_template(request)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()

        self.assertEqual(
            rows,
            [(
                "S PNR",
                "Client ID",
                "Ticket No",
                "Passenger Name",
                "Sector",
                "Booking Type",
                "Travel Type",
                "Airline Name",
                "Airline Type",
                "Airline PNR",
                "Net Amount",
                "Travel Start Date",
                "Travel End Date",
                "Booking Date",
                "Status",
            )],
        )


class CouponAvailabilityTimingTests(SimpleTestCase):
    @patch("core.views.timezone.localdate", return_value=date(2026, 10, 8))
    def test_coupon_availability_timing_rules(self, _localdate):
        # Current date: 2026-10-08
        test_cases = [
            ("2026-10-07", "EXPIRED", "Cancelled booking -> coupon must be void"),
            ("2026-10-08", "PENDING", "Travel Start Date is today -> must remain PENDING"),
            ("2026-10-09", "PENDING", "Travel Start Date is tomorrow -> must be PENDING"),
            ("2026-10-15", "PENDING", "Travel Start Date is in future -> must be PENDING"),
            ("2026-10-07", "AVAILABLE", "Travel Start Date is yesterday -> must be AVAILABLE"),
            ("2026-10-06", "AVAILABLE", "Travel Start Date is 2 days ago -> must be AVAILABLE"),
        ]

        for travel_date_str, expected_status, description in test_cases:
            coupon = {
                "booking_id": "BK-1",
                "travel_start_date": travel_date_str,
                "status": "PENDING",
            }
            bookings = [
                {
                    "booking_id": "BK-1",
                    "travel_date": travel_date_str,
                    "status": "Cancelled" if expected_status == "EXPIRED" else "Confirmed",
                }
            ]
            actual_status = _coupon_status_for_travel(coupon, bookings)
            self.assertEqual(
                actual_status,
                expected_status,
                f"Failed for {travel_date_str}: {description} (got {actual_status}, expected {expected_status})",
            )

    @patch("core.views.timezone.localdate", return_value=date(2026, 10, 8))
    @patch("core.views.write")
    @patch("core.views.read")
    def test_auto_release_coupons_updates_only_past_travel_dates(self, mock_read, mock_write, _localdate):
        # Today is 2026-10-08
        mock_read.side_effect = lambda key: {
            "coupons": [
                {
                    "booking_id": "BK-TODAY",
                    "travel_start_date": "2026-10-08",
                    "status": "PENDING",
                    "points": Decimal("10.00"),
                    "is_released": False,
                },
                {
                    "booking_id": "BK-TOMORROW",
                    "travel_start_date": "2026-10-09",
                    "status": "PENDING",
                    "points": Decimal("20.00"),
                    "is_released": False,
                },
                {
                    "booking_id": "BK-YESTERDAY",
                    "travel_start_date": "2026-10-07",
                    "status": "PENDING",
                    "points": Decimal("30.00"),
                    "is_released": False,
                },
                {
                    "booking_id": "BK-PAST-2DAYS",
                    "travel_start_date": "2026-10-06",
                    "status": "PENDING",
                    "points": Decimal("40.00"),
                    "is_released": False,
                },
            ],
            "bookings": [
                {"booking_id": "BK-TODAY", "travel_date": "2026-10-08", "net_amount": Decimal("1000")},
                {"booking_id": "BK-TOMORROW", "travel_date": "2026-10-09", "net_amount": Decimal("2000")},
                {"booking_id": "BK-YESTERDAY", "travel_date": "2026-10-07", "net_amount": Decimal("3000")},
                {"booking_id": "BK-PAST-2DAYS", "travel_date": "2026-10-06", "net_amount": Decimal("4000")},
            ],
        }[key]

        from core.views import auto_release_coupons
        updated_coupons = auto_release_coupons()

        coupon_status_map = {c["booking_id"]: (c["status"], c["is_released"]) for c in updated_coupons}

        self.assertEqual(coupon_status_map["BK-TODAY"], ("PENDING", False))
        self.assertEqual(coupon_status_map["BK-TOMORROW"], ("PENDING", False))
        self.assertEqual(coupon_status_map["BK-YESTERDAY"], ("AVAILABLE", True))
        self.assertEqual(coupon_status_map["BK-PAST-2DAYS"], ("AVAILABLE", True))

    def test_calculate_discounted_net_negative_amount(self):
        from core.views import calculate_discounted_net
        amount, pct, rule = calculate_discounted_net("-450.50", "Air India", "Economy Class")
        self.assertEqual(amount, Decimal("-450.50"))
        self.assertEqual(pct, Decimal("0"))
        self.assertIsNone(rule)


